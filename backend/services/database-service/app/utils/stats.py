"""Candidate analytics / funnel helpers."""
from __future__ import annotations

from datetime import datetime, timezone
from typing import Optional

from openpyxl import Workbook
import tempfile
from sqlalchemy import select
from sqlalchemy.orm import selectinload
from sqlalchemy.ext.asyncio import AsyncSession

from app.db.v1.enums import ARCHIVE_STATUSES, CandidateStage, CandidateStatus
from app.db.v1.models import Candidate, CandidateVacancyRelation


FUNNEL_ORDER: list[str] = [s.value for s in CandidateStatus]


def empty_funnel_counts() -> dict[str, int]:
    return {status: 0 for status in FUNNEL_ORDER}


def _naive_utc(value: datetime | None) -> datetime | None:
    """Normalize timestamps so naive export bounds can be compared with aware DB values."""
    if value is None:
        return None
    if getattr(value, "tzinfo", None) is not None:
        return value.astimezone(timezone.utc).replace(tzinfo=None)
    return value


def _relation_sort_key(rel: CandidateVacancyRelation):
    ts = _naive_utc(rel.updated_at or rel.created_at) or datetime.min
    return (ts, rel.id or 0)


def resolve_candidate_funnel_status(candidate: Candidate) -> Optional[str]:
    """
    Map a candidate to one funnel status label.

    Priority:
    - hired stage → ВНР (or уволился if last relation says so)
    - blacklist → last relation status or «не подходит»
    - archive → last relation status or «отказ»
    - employment → active relation, else latest relation
    """
    relations = sorted(
        list(candidate.vacancies or []),
        key=_relation_sort_key,
        reverse=True,
    )
    active = next((r for r in relations if r.is_active), None)
    latest = relations[0] if relations else None

    if candidate.stage == CandidateStage.hired:
        if latest and latest.status == CandidateStatus.resigned:
            return CandidateStatus.resigned.value
        return CandidateStatus.started_work.value

    if candidate.stage == CandidateStage.blacklisted:
        if latest and latest.status is not None:
            return latest.status.value if hasattr(latest.status, "value") else str(latest.status)
        return CandidateStatus.not_suitable.value

    if candidate.stage == CandidateStage.archieved:
        if latest and latest.status is not None:
            val = latest.status.value if hasattr(latest.status, "value") else str(latest.status)
            # Prefer precise archive funnel labels; fall back to «отказ»
            if val in {s.value for s in ARCHIVE_STATUSES} or val in FUNNEL_ORDER:
                return val
        return CandidateStatus.rejection.value

    # in hiring / employment
    pick = active or latest
    if pick and pick.status is not None:
        return pick.status.value if hasattr(pick.status, "value") else str(pick.status)
    return None


def _candidates_with_relations_stmt(vacancy_id: int | None = None):
    """Load candidates and their vacancy relations.

    Filter by vacancy via IN (candidate_id), not JOIN + DISTINCT: PostgreSQL
    has no equality operator for json columns on candidates, so
    SELECT DISTINCT candidates.* fails with UndefinedFunctionError.
    """
    stmt = select(Candidate).options(
        selectinload(Candidate.vacancies).selectinload(CandidateVacancyRelation.vacancy)
    )
    if vacancy_id is not None:
        related_ids = select(CandidateVacancyRelation.candidate_id).where(
            CandidateVacancyRelation.vacancy_id == int(vacancy_id)
        )
        stmt = stmt.where(Candidate.id.in_(related_ids))
    return stmt


def resolve_candidate_funnel_status_for_vacancy(
    candidate: Candidate,
    vacancy_id: int,
) -> Optional[str]:
    """Resolve funnel status for a concrete vacancy relation."""
    relations = [
        rel
        for rel in (candidate.vacancies or [])
        if int(getattr(rel, "vacancy_id", 0) or 0) == int(vacancy_id)
    ]
    if not relations:
        return None
    relations = sorted(relations, key=_relation_sort_key, reverse=True)
    pick = relations[0]
    if pick and pick.status is not None:
        return pick.status.value if hasattr(pick.status, "value") else str(pick.status)
    return None


async def get_candidate_status_summary(
    db: AsyncSession,
    vacancy_id: int | None = None,
) -> dict[str, int]:
    """Count every candidate once under the new funnel statuses."""
    summary = empty_funnel_counts()
    result = await db.execute(_candidates_with_relations_stmt(vacancy_id))
    for candidate in result.scalars().all():
        status = (
            resolve_candidate_funnel_status_for_vacancy(candidate, vacancy_id)
            if vacancy_id is not None
            else resolve_candidate_funnel_status(candidate)
        )
        if not status:
            continue
        if status not in summary:
            summary[status] = 0
        summary[status] += 1
    return summary


def build_funnel_xlsx(
    funnel_data: dict[str, int],
    detail_data: dict[str, list[tuple[str, str]]],
    date_from: datetime | None,
    date_to: datetime | None,
) -> str:
    wb = Workbook()
    ws_summary = wb.active
    ws_summary.title = "Воронка"

    ws_summary.append(["Статус", "Количество кандидатов"])
    for status in FUNNEL_ORDER:
        ws_summary.append([status, int(funnel_data.get(status, 0))])
    # any unexpected leftovers
    for status, count in sorted(funnel_data.items()):
        if status not in FUNNEL_ORDER:
            ws_summary.append([status, count])

    for status in FUNNEL_ORDER:
        candidates_list = detail_data.get(status) or []
        if not candidates_list:
            continue
        ws = wb.create_sheet(title=status[:31])
        ws.append(["ФИО кандидата", "Вакансия"])
        for full_name, vacancy_name in candidates_list:
            ws.append([full_name, vacancy_name])

    for status, candidates_list in detail_data.items():
        if status in FUNNEL_ORDER:
            continue
        ws = wb.create_sheet(title=str(status)[:31])
        ws.append(["ФИО кандидата", "Вакансия"])
        for full_name, vacancy_name in candidates_list:
            ws.append([full_name, vacancy_name])

    temp_file = tempfile.NamedTemporaryFile(delete=False, suffix=".xlsx")
    wb.save(temp_file.name)
    return temp_file.name


def _vacancy_name_for_candidate(candidate: Candidate, vacancy_id: int | None = None) -> str:
    relations = list(candidate.vacancies or [])
    if vacancy_id is not None:
        targeted = [
            rel
            for rel in relations
            if int(getattr(rel, "vacancy_id", 0) or 0) == int(vacancy_id)
        ]
        if targeted:
            relations = targeted
    relations = sorted(relations, key=_relation_sort_key, reverse=True)
    for rel in relations:
        if rel.vacancy and getattr(rel.vacancy, "name", None):
            return rel.vacancy.name
    return "—"


async def get_funnel_data(
    db: AsyncSession,
    date_from: datetime | None,
    date_to: datetime | None,
    vacancy_id: int | None = None,
):
    """
    Funnel snapshot for export.
    When a date range is set, include candidates whose own updated_at
    or any vacancy-relation updated_at falls in the range.
    """
    result = await db.execute(_candidates_with_relations_stmt(vacancy_id))
    candidates: list[Candidate] = list(result.scalars().all())

    funnel_data = empty_funnel_counts()
    detail_data: dict[str, list[tuple[str, str]]] = {s: [] for s in FUNNEL_ORDER}

    for candidate in candidates:
        if date_from and date_to:
            bound_from = _naive_utc(date_from)
            bound_to = _naive_utc(date_to)
            cand_ts = _naive_utc(candidate.updated_at or candidate.created_at)
            in_range = bool(cand_ts and bound_from and bound_to and bound_from <= cand_ts <= bound_to)
            if not in_range:
                for rel in candidate.vacancies or []:
                    rel_ts = _naive_utc(rel.updated_at or rel.created_at)
                    if rel_ts and bound_from and bound_to and bound_from <= rel_ts <= bound_to:
                        in_range = True
                        break
            if not in_range:
                continue

        status = (
            resolve_candidate_funnel_status_for_vacancy(candidate, vacancy_id)
            if vacancy_id is not None
            else resolve_candidate_funnel_status(candidate)
        )
        if not status:
            continue
        funnel_data[status] = funnel_data.get(status, 0) + 1
        detail_data.setdefault(status, []).append(
            (candidate.full_name, _vacancy_name_for_candidate(candidate, vacancy_id))
        )

    return funnel_data, detail_data
