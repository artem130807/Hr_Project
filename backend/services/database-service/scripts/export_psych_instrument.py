# -*- coding: utf-8 -*-
"""Export complex work-profile instrument from new_test_info xlsx into JSON."""
from __future__ import annotations

import json
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[4]
XLSX = next(Path(ROOT / "new_test_info").glob("*.xlsx"))
OUT = (
    Path(__file__).resolve().parents[1]
    / "app"
    / "psychometrics"
    / "instruments"
    / "complex_work_profile.json"
)

DISC_CONTEXTS = {
    "В работе": "В обычной рабочей ситуации",
    "Под давлением": "Когда сроки сжаты или давление высокое",
    "Личное": "В привычном для вас способе действовать",
}

ASPECTS = [
    {"code": "Ee", "name": "Энтузиазм", "factor": "E"},
    {"code": "Ea", "name": "Ассертивность", "factor": "E"},
    {"code": "Ac", "name": "Сострадание", "factor": "A"},
    {"code": "Ap", "name": "Вежливость", "factor": "A"},
    {"code": "Ci", "name": "Трудолюбие", "factor": "C"},
    {"code": "Co", "name": "Прилежность", "factor": "C"},
    {"code": "Nv", "name": "Волатильность", "factor": "N"},
    {"code": "Nw", "name": "Раздражительность", "factor": "N"},
    {"code": "Oi", "name": "Интеллект", "factor": "O"},
    {"code": "Oo", "name": "Открытость", "factor": "O"},
]

FACTORS = [
    {"code": "E", "name": "Экстраверсия", "aspects": ["Ee", "Ea"], "invert": False},
    {"code": "A", "name": "Доброжелательность", "aspects": ["Ac", "Ap"], "invert": False},
    {"code": "C", "name": "Добросовестность", "aspects": ["Ci", "Co"], "invert": False},
    {"code": "ES", "name": "Эмоциональная стабильность", "aspects": ["Nv", "Nw"], "invert": True},
    {"code": "O", "name": "Открытость и интеллект", "aspects": ["Oi", "Oo"], "invert": False},
]

PAEI = [
    {"code": "P", "name": "Производитель результата"},
    {"code": "A", "name": "Администратор"},
    {"code": "E", "name": "Предприниматель"},
    {"code": "I", "name": "Интегратор"},
]

DISC_META = [
    {"code": "D", "name": "D"},
    {"code": "I", "name": "I"},
    {"code": "S", "name": "S"},
    {"code": "C", "name": "C"},
]


def _cell(ws, r, c):
    v = ws.cell(r, c).value
    if v is None:
        return None
    if isinstance(v, str):
        return v.strip()
    return v


def main():
    wb = load_workbook(XLSX, data_only=False)
    keys = wb["Ключи"]
    settings = wb["Настройки"]

    thresholds = {}
    for row in range(9, 15):
        param = _cell(settings, row, 1)
        value = _cell(settings, row, 2)
        if param:
            thresholds[str(param)] = value

    items = []
    num = 0
    for row in range(5, 133):
        item_id = _cell(keys, row, 1)
        module = _cell(keys, row, 2)
        context = _cell(keys, row, 3)
        prompt = _cell(keys, row, 4)
        if not item_id:
            continue
        num += 1
        opts = {
            "A": _cell(keys, row, 5),
            "B": _cell(keys, row, 6),
            "C": _cell(keys, row, 7),
            "D": _cell(keys, row, 8),
        }
        cats = {
            "A": _cell(keys, row, 9),
            "B": _cell(keys, row, 10),
            "C": _cell(keys, row, 11),
            "D": _cell(keys, row, 12),
        }
        reverse = _cell(keys, row, 13)
        factor = _cell(keys, row, 14)
        scores = {
            "A": _cell(keys, row, 15),
            "B": _cell(keys, row, 16),
            "C": _cell(keys, row, 17),
            "D": _cell(keys, row, 18),
        }

        if module == "DISC":
            items.append(
                {
                    "num": num,
                    "code": str(item_id),
                    "module": "disc",
                    "context": str(context or ""),
                    "prompt": str(prompt or DISC_CONTEXTS.get(str(context or ""), "")),
                    "options": {k: str(v or "") for k, v in opts.items()},
                    "categories": {k: str(v or "") for k, v in cats.items()},
                }
            )
        elif module == "AVP":
            aspect_name = str(context or "")
            aspect_code = next((a["code"] for a in ASPECTS if a["name"] == aspect_name), "")
            items.append(
                {
                    "num": num,
                    "code": str(item_id),
                    "module": "avp",
                    "aspect": aspect_name,
                    "aspect_code": aspect_code,
                    "factor": str(factor or ""),
                    "text": str(prompt or ""),
                    "reverse": int(reverse or 0),
                }
            )
        elif module == "SJT":
            items.append(
                {
                    "num": num,
                    "code": str(item_id),
                    "module": "sjt",
                    "title": str(context or ""),
                    "prompt": str(prompt or ""),
                    "options": {k: str(v or "") for k, v in opts.items()},
                    "priorities": {k: str(v or "") for k, v in cats.items()},
                    "scores": {
                        k: int(v) if v not in (None, "") else 0 for k, v in scores.items()
                    },
                }
            )

    payload = {
        "id": "complex_work_profile",
        "aliases": ["complex_work_behavior_220"],
        "version": "1.0",
        "assessment_version": "1.0",
        "item_bank_version": "1.0",
        "scoring_version": "1.0-pilot",
        "norms_version": "none",
        "name": "Комплексная оценка рабочего профиля",
        "description": (
            "128 заданий: рабочее поведение, утверждения о привычном стиле "
            "и рабочие ситуации. Отвечайте о своём обычном поведении."
        ),
        "duration_minutes": "25–40",
        "answer_scale": {
            "min": 1,
            "max": 5,
            "labels": {
                "1": "совсем не похоже",
                "2": "скорее не похоже",
                "3": "отчасти похоже",
                "4": "скорее похоже",
                "5": "очень похоже",
            },
        },
        "participant_fields": [
            {"id": "full_name", "label": "ФИО", "required": True},
            {"id": "position", "label": "Должность / вакансия", "required": True},
            {"id": "birth_date", "label": "Дата рождения", "required": True, "type": "date"},
            {"id": "taken_at", "label": "Дата прохождения", "required": True, "type": "date", "default": "today"},
        ],
        "instruction": (
            "Отвечайте о своем обычном поведении. Здесь нет идеального профиля. "
            "Выбирайте вариант, который точнее описывает вас или ваше вероятное действие, "
            "а не тот, который кажется наиболее правильным."
        ),
        "presentation": {
            "scheme": "231",
            "block_order": ["avp", "sjt", "disc"],
            "shuffle_modules": ["avp"],
            "blocks": [
                {"module": "avp", "title": "Привычный стиль"},
                {"module": "sjt", "title": "Рабочие ситуации"},
                {"module": "disc", "title": "Рабочее поведение"},
            ],
        },
        "items": items,
        "avp_aspects": ASPECTS,
        "avp_factors": FACTORS,
        "disc": DISC_META,
        "paei": PAEI,
        "quality_thresholds": {
            "critical_speed_sec": float(thresholds.get("critical_speed_sec") or 480),
            "warning_speed_sec": float(thresholds.get("warning_speed_sec") or 900),
            "rapid_answer_ms": float(thresholds.get("rapid_answer_ms") or 1500),
            "rapid_share_warning": float(thresholds.get("rapid_share_warning") or 0.25),
            "avp_low_variance_sd": float(thresholds.get("avp_low_variance_sd") or 0.25),
            "inactivity_timeout_sec": float(thresholds.get("inactivity_timeout_sec") or 60),
            "one_minute_wall_sec": 70,
            "expected_items": 128,
        },
    }

    OUT.parent.mkdir(parents=True, exist_ok=True)
    OUT.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")
    counts = {}
    for it in items:
        counts[it["module"]] = counts.get(it["module"], 0) + 1
    print(f"Wrote {OUT} items={len(items)} {counts}")


if __name__ == "__main__":
    main()
